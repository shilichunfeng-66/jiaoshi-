# -*- coding: utf-8 -*-
"""将3张科研图插入护膝申报书（修正版）"""
import shutil
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

SRC = r'D:/科研文件/outputs/池州学院_大创申报书_护膝.docx'
DST = r'D:/科研文件/outputs/池州学院_大创申报书_护膝_v3.docx'
shutil.copy2(SRC, DST)

doc = Document(DST)
cell = doc.tables[3].cell(0, 0)
tc = cell._tc

PIC = {
    '图1': (r'D:/科研文件/outputs/fig_护膝_制备工艺.png', '图1 植物纤维混杂复合材料制备工艺流程'),
    '图2': (r'D:/科研文件/outputs/fig_护膝_铺层结构.png', '图2 混杂复合材料两种铺层结构示意'),
    '图3': (r'D:/科研文件/outputs/fig_护膝_配比性能.png', '图3 混杂配比与力学性能关系（正混杂效应示意）'),
}

def make_cap_para(caption):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr'); p.append(pPr)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); pPr.append(jc)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:line'), '360'); sp.set(qn('w:lineRule'), 'auto'); pPr.append(sp)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:eastAsia'), '宋体'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rf)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '21'); rPr.append(sz)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), '21'); rPr.append(szcs)
    r.append(rPr)
    t = OxmlElement('w:t'); t.text = caption; r.append(t)
    p.append(r)
    return p

def is_heading(txt):
    return bool(re.match(r'^[（(][一二三四五六][）)]', txt)) or bool(re.match(r'^[1-9]\.', txt)) or bool(re.match(r'^第[一二三四]阶段', txt)) or bool(re.match(r'^[一二三四五六]、', txt))

def insert_pic_after_title(cell, anchor_title, pic_path, caption, width_cm=11):
    all_p = list(tc.findall(qn('w:p')))
    target_idx = None
    for i, p in enumerate(all_p):
        txt = ''.join(t.text or '' for t in p.findall('.//' + qn('w:t')))
        if txt.strip().startswith(anchor_title):
            target_idx = i; break
    if target_idx is None:
        print(f'  未找到标题: {anchor_title}')
        return False
    next_heading = None
    for j in range(target_idx+1, len(all_p)):
        txt = ''.join(t.text or '' for t in all_p[j].findall('.//' + qn('w:t')))
        if is_heading(txt):
            next_heading = all_p[j]; break
    if next_heading is None:
        print(f'  无下一标题'); return False
    # 先插图注（紧贴 next_heading 之前）
    next_heading.addprevious(make_cap_para(caption))
    # 创建图片段并移到图注前（保证图片在上，图注在下）
    pic_p = cell.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(pic_path, width=Cm(width_cm))
    cap_p = next_heading.getprevious()
    if cap_p is not None and cap_p.tag == qn('w:p'):
        cap_p.addprevious(pic_p._p)
    print(f'  ✓ {caption[:18]}... -> {anchor_title[:16]}... 后')
    return True

# 插入顺序：先插图3（"4."后），再插图1和图2（都在"2."后）
# 图2 在图1 之后插入（保持物理顺序：图1 -> 图2 -> 图3）
print('=== 先插图3 ===')
insert_pic_after_title(cell, '4. 混杂配比与性能优化', PIC['图3'][0], PIC['图3'][1])
print('=== 插图1 ===')
insert_pic_after_title(cell, '2. 混杂复合材料的制备工艺', PIC['图1'][0], PIC['图1'][1])
print('=== 插图2 ===')
insert_pic_after_title(cell, '2. 混杂复合材料的制备工艺', PIC['图2'][0], PIC['图2'][1])

doc.save(DST)
print(f'\n保存: {DST}')

doc2 = Document(DST)
c = doc2.tables[3].cell(0, 0)
print(f'\n=== 物理顺序验证 ===')
for i, p in enumerate(c.paragraphs):
    xml = p._p.xml
    has_pic = 'pic:pic' in xml or 'blip' in xml
    txt = p.text.strip()
    if has_pic:
        print(f'  P[{i:2d}] [图片]')
    elif txt.startswith('图') and len(txt) < 30:
        print(f'  P[{i:2d}] [图注] {txt}')
print(f'\n=== 图片总数: {len(doc2.inline_shapes)} ===')