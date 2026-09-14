# -*- coding: utf-8 -*-
"""健美操申报书 — 组装写入新 docx（不配图）
基于池州学院母版，清空表3行0，写入正文，自动填项目名称到表2行0列1。
"""
import shutil, re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r'D:/科研文件/schools/池州学院/大创/母版_申报书.docx'
DST = r'D:/科研文件/outputs/池州学院_大创申报书_健美操.docx'
CONTENT_TXT = r'D:/科研文件/outputs/_正文_健美操.txt'
TOPIC = '全民健身战略下大众健美操标准化推广路径研究'

shutil.copy2(SRC, DST)
doc = Document(DST)
cell = doc.tables[3].cell(0, 0)
tc = cell._tc

# 清空表3行0
for p in list(tc.findall(qn('w:p'))):
    tc.remove(p)

FONT = '仿宋'
SZ_H1 = 28    # 四号 14pt
SZ_BODY = 24  # 小四 12pt

def _set_run_fonts(rPr, font, size_half, bold=False):
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:eastAsia'), font)
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rf)
    if bold:
        rPr.append(OxmlElement('w:b'))
        rPr.append(OxmlElement('w:bCs'))
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(size_half)); rPr.append(sz)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(size_half)); rPr.append(szcs)

def _set_para_format(pPr, indent=True, line=360, space_before=120, space_after=120):
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:line'), str(line)); sp.set(qn('w:lineRule'), 'auto')
    sp.set(qn('w:before'), str(space_before)); sp.set(qn('w:after'), str(space_after))
    pPr.append(sp)
    if indent:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLineChars'), '200')
        ind.set(qn('w:firstLine'), '480')
        pPr.append(ind)

def make_para(text, kind):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr'); p.append(pPr)
    if kind == 'h1':
        _set_para_format(pPr, indent=False, line=360, space_before=160, space_after=100)
        r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
        _set_run_fonts(rPr, FONT, SZ_H1, bold=True)
        r.append(rPr)
        t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
        p.append(r)
    elif kind == 'subhead':
        _set_para_format(pPr, indent=True, line=360, space_before=80, space_after=80)
        m = re.match(r'^(\d+[\.、]\s*[^：:]+[：:]\s*)(.*)$', text)
        if m:
            bold_part, rest = m.group(1), m.group(2)
            r1 = OxmlElement('w:r'); rPr1 = OxmlElement('w:rPr')
            _set_run_fonts(rPr1, FONT, SZ_BODY, bold=True)
            r1.append(rPr1)
            t1 = OxmlElement('w:t'); t1.text = bold_part; t1.set(qn('xml:space'), 'preserve'); r1.append(t1)
            p.append(r1)
            if rest:
                r2 = OxmlElement('w:r'); rPr2 = OxmlElement('w:rPr')
                _set_run_fonts(rPr2, FONT, SZ_BODY, bold=False)
                r2.append(rPr2)
                t2 = OxmlElement('w:t'); t2.text = rest; t2.set(qn('xml:space'), 'preserve'); r2.append(t2)
                p.append(r2)
        else:
            r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
            _set_run_fonts(rPr, FONT, SZ_BODY, bold=False)
            r.append(rPr)
            t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
            p.append(r)
    else:
        _set_para_format(pPr, indent=True, line=360, space_before=100, space_after=100)
        bold = (kind in ('h2', 'h3', 'stage'))
        r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
        _set_run_fonts(rPr, FONT, SZ_BODY, bold=bold)
        r.append(rPr)
        t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
        p.append(r)
    return p

def classify(text):
    t = text.strip()
    if not t: return None
    if re.match(r'^[一二三四五六七八九十]+、', t): return 'h1'
    if re.match(r'^[（(][一二三四五六七八九十]+[）)]', t): return 'h2'
    if re.match(r'^第[一二三四五六七八九十]+阶段', t): return 'stage'
    if re.match(r'^\d+[\.、]\s*\S+[：:]', t): return 'subhead'
    if re.match(r'^[1-9][\.、]', t): return 'h3'
    return 'body'

# 写入正文
with open(CONTENT_TXT, encoding='utf-8') as f:
    lines = f.read().split('\n')

count = 0
for line in lines:
    t = line.strip()
    if not t: continue
    kind = classify(t)
    if kind is None: continue
    tc.append(make_para(t, kind))
    count += 1

print(f'正文写入 {count} 段')

# 填项目名称到表2行0列1
if TOPIC:
    t2 = doc.tables[2]
    tr2 = t2._tbl.findall(qn('w:tr'))[0]
    tcs2 = tr2.findall(qn('w:tc'))
    if len(tcs2) >= 2:
        cell_topic = tcs2[1]
        ps = cell_topic.findall(qn('w:p'))
        if ps:
            p_topic = ps[0]
            r_topic = OxmlElement('w:r')
            rPr_topic = OxmlElement('w:rPr')
            rf_topic = OxmlElement('w:rFonts')
            rf_topic.set(qn('w:eastAsia'), '宋体')
            rf_topic.set(qn('w:ascii'), 'Times New Roman')
            rf_topic.set(qn('w:hAnsi'), 'Times New Roman')
            rPr_topic.append(rf_topic)
            rPr_topic.append(OxmlElement('w:b'))
            rPr_topic.append(OxmlElement('w:bCs'))
            sz_topic = OxmlElement('w:sz'); sz_topic.set(qn('w:val'), '28'); rPr_topic.append(sz_topic)
            szcs_topic = OxmlElement('w:szCs'); szcs_topic.set(qn('w:val'), '28'); rPr_topic.append(szcs_topic)
            r_topic.append(rPr_topic)
            t_topic = OxmlElement('w:t'); t_topic.text = TOPIC; t_topic.set(qn('xml:space'), 'preserve'); r_topic.append(t_topic)
            p_topic.append(r_topic)
            print(f'已填项目名称: {TOPIC}')

doc.save(DST)
print(f'保存: {DST}')

# 验证
doc2 = Document(DST)
c = doc2.tables[3].cell(0, 0)
total = sum(len(p.text) for p in c.paragraphs)
print(f'验证：表3行0 段落数 {len(c.paragraphs)}，总字数 {total}')
