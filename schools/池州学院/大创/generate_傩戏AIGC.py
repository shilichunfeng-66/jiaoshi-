# -*- coding: utf-8 -*-
"""AIGC赋能池州傩戏的短视频传播与数字化传承研究 — 池州学院申报书生成"""
import os, shutil
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = r'D:/科研文件/schools/池州学院/大创'
MASTER = os.path.join(HERE, '母版_申报书.docx')
OUTPUT = r'D:/科研文件/outputs/池州学院_大创申报书_傩戏AIGC.docx'
FONT = '仿宋'

def new_para(size=12, bold=False, indent=True):
    """创建段落XML"""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr'); p.append(pPr)
    # 行距1.5
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:line'), '360'); sp.set(qn('w:lineRule'), 'auto')
    sp.set(qn('w:before'), '120'); sp.set(qn('w:after'), '120')
    pPr.append(sp)
    # 首行缩进2字符
    if indent:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLineChars'), '200'); ind.set(qn('w:firstLine'), '480')
        pPr.append(ind)
    # run
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:eastAsia'), FONT); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rf)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(size*2)); rPr.append(sz)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(size*2)); rPr.append(szcs)
    if bold:
        rPr.append(OxmlElement('w:b'))
    r.append(rPr)
    p.append(r)
    return p

def set_text(p, text):
    t = p.find(qn('w:r')).find(qn('w:t'))
    if t is None:
        t = OxmlElement('w:t')
        p.find(qn('w:r')).append(t)
    t.text = text
    t.set(qn('xml:space'), 'preserve')

def title_para(text, size=12, bold=True):
    p = new_para(size=size, bold=bold, indent=False)
    set_text(p, text)
    return p

def body_para(text):
    p = new_para(size=12, bold=False, indent=True)
    set_text(p, text)
    return p

def subhead_para(title, body):
    """分点段落: 编号+小标题加粗,描述不加粗"""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr'); p.append(pPr)
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:line'), '360'); sp.set(qn('w:lineRule'), 'auto')
    sp.set(qn('w:before'), '120'); sp.set(qn('w:after'), '120')
    pPr.append(sp)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:firstLineChars'), '200'); ind.set(qn('w:firstLine'), '480')
    pPr.append(ind)
    # 加粗标题run
    r1 = OxmlElement('w:r')
    rPr1 = OxmlElement('w:rPr')
    rf1 = OxmlElement('w:rFonts')
    rf1.set(qn('w:eastAsia'), FONT); rf1.set(qn('w:ascii'), 'Times New Roman'); rf1.set(qn('w:hAnsi'), 'Times New Roman')
    rPr1.append(rf1)
    sz1 = OxmlElement('w:sz'); sz1.set(qn('w:val'), '24'); rPr1.append(sz1)
    rPr1.append(OxmlElement('w:b'))
    r1.append(rPr1)
    t1 = OxmlElement('w:t'); t1.text = title; r1.append(t1)
    p.append(r1)
    # 不加粗描述run
    r2 = OxmlElement('w:r')
    rPr2 = OxmlElement('w:rPr')
    rf2 = OxmlElement('w:rFonts')
    rf2.set(qn('w:eastAsia'), FONT); rf2.set(qn('w:ascii'), 'Times New Roman'); rf2.set(qn('w:hAnsi'), 'Times New Roman')
    rPr2.append(rf2)
    sz2 = OxmlElement('w:sz'); sz2.set(qn('w:val'), '24'); rPr2.append(sz2)
    r2.append(rPr2)
    t2 = OxmlElement('w:t'); t2.text = body; r2.append(t2)
    p.append(r2)
    return p

shutil.copy2(MASTER, OUTPUT)
doc = Document(OUTPUT)
t3 = doc.tables[3]

# ===== 重建表3 行0(正文一~六章) =====
cell = t3.cell(0, 0)
tc = cell._tc
# 清空所有段落
for p in tc.findall(qn('w:p')):
    tc.remove(p)

def add_title(text, size=12):
    p = title_para(text, size=size)
    tc.append(p)

def add_body(text):
    tc.append(body_para(text))

def add_sub(title, body):
    tc.append(subhead_para(title, body))

# ===== 一、项目实施的目的、意义 =====
add_title('一、项目实施的目的、意义', size=14)

add_title('（一）项目简介')
add_body('池州傩戏被誉为"戏曲活化石"，2006年入选首批国家级非物质文化遗产名录，以佩戴彩绘面具表演为核心，'
    '融合歌舞、戏曲、民俗与祭祀仪式，承载着皖南宗族社会的请神敬祖、驱邪纳福传统。然而长期以来，傩戏的'
    '传播高度依赖线下演出与展馆陈列，受时空限制明显，年轻受众触达率偏低。本项目尝试用AIGC技术破解这一'
    '困境：以人工智能生成内容（AIGC）为工具，把傩面具纹样、傩舞动作、傩仪场景转化为适配短视频平台的'
    '数字内容，探索"传统非遗+新技术+新传播"的活态传承路径。')
add_body('项目立足池州本地资源优势，依托AIGC文生图、文生视频、数字人等工具，围绕"文化梳理—内容生产—'
    '平台传播—资源沉淀"四个环节展开：先系统梳理池州傩戏的文化基因，再用AIGC重构傩面具与傩舞的视觉'
    '表达，进而制作系列短视频在多平台发布，最终沉淀为可复用的傩戏数字资源库，为池州傩戏的数字化传承'
    '留下一套可持续的内容生产范式。')

add_title('（二）研究实施目的')
add_body('研究的目的，首先是回答一个现实问题：傩戏这门古老的艺术，怎样借助AIGC让年轻人重新看见它。'
    '池州傩戏的传播困境是具体的——它的美藏在傩面具的狞厉纹样里、藏在傩舞的刚劲程式里，但这些内容'
    '若不经过适合短视频的转化，就很难进入年轻人的信息流。现有传播多停留在"记录式"拍摄，把傩戏演出'
    '原样搬上屏幕，热闹过后难以留下持续关注。本项目的目的，是用AIGC把傩戏的文化元素"拆解再重组"，'
    '生成既有文化辨识度、又符合短视频审美的内容形态。')
add_body('其次，研究希望探索一条可复制的技术路径。傩戏不是孤例，青阳腔、东至花灯等池州其他非遗同样'
    '面临传播难题。如果本项目能跑通"傩戏—AIGC—短视频"的内容生产流程，形成一套提示词模板、素材库'
    '和制作规范，这套方法就可以推广到其他非遗项目，从"单点试水"走向"体系复制"。这是研究在方法层面'
    '的价值。')
add_body('最后，研究还有一个朴素的愿望：让技术真正服务于本地。池州学院的学生守着傩戏这一文化富矿，'
    '如果只是看别人用AI做非遗，自己不动手，未免可惜。通过这个项目，把AI工具用得起来、把傩戏内容'
    '做得出来，既是专业训练，也是对家乡文化的一份贡献。')

add_title('（三）研究实施意义')
add_body('从文化传承看，项目的意义在于为傩戏的"活态传承"补上数字这一环。傩戏的传承靠的是宗族里的'
    '口传心授、年复一年的傩事活动，这是它的根；但要让根长得更开，需要数字内容把傩戏的样貌带到更远'
    '的地方。AIGC的价值恰恰在于降低内容生产门槛——过去做一支傩戏动画需要专业团队数月打磨，如今借助'
    'AI工具，学生也能在较短时间内生成可用的视觉素材。这让"人人可参与非遗传播"有了技术可能。')
add_body('从技术应用看，项目的意义在于回答"AI做非遗，边界在哪里"。傩戏的文化本真性很强，傩面具的'
    '每一处纹样、傩舞的每一个程式都有民俗依据。AIGC生成若不加约束，容易产出"长得像傩、实则走样"的'
    '伪非遗内容。本项目把"传承人把关、文化馆校准"作为硬性流程，探索AI创新与守正的平衡点，这正是'
    '当前非遗数字化领域普遍关心、又缺乏实践答案的问题。')
add_body('从地方发展看，项目的意义在于为池州文旅贡献可用的传播素材。池州正在打造长三角休闲度假'
    '旅游目的地，傩戏是最具辨识度的文化名片之一。项目产出的短视频内容、数字资源库，可以作为地方'
    '文旅宣传的轻量化素材，降低文旅部门的内容生产成本，实现"学生实践—文化传播—地方发展"的良性循环。')

# ===== 二、项目研究内容和拟解决的关键问题 =====
add_title('二、项目研究内容和拟解决的关键问题', size=14)

add_title('（一）项目研究内容')
add_body('研究内容围绕"梳理—生产—传播—沉淀"四个环节展开，形成环环相扣的内容体系。')

add_title('1. 池州傩戏文化资源的系统梳理')
add_body('这是研究的基础环节。傩戏的文化信息分散在傩面具、傩舞、傩仪、唱本、宗族谱牒等多个载体中，'
    '需要先做系统整理。研究将以池州傩文化展示馆、梅街镇刘街社区等傩戏发源地为调研对象，梳理傩戏的'
    '历史渊源、面具谱系、经典剧目和民俗内涵，重点提炼12尊核心傩面具角色（和仙、观音、寿星、鲍三娘、'
    '周仓、二郎神、城隍、钟馗、千里眼、顺风耳等）的纹样特征与文化寓意，建立傩戏文化基因的基础档案，'
    '为后续的AIGC创作提供"有据可依"的素材来源。')

add_title('2. 傩戏AIGC短视频的内容生产模式')
add_body('这是研究的核心环节。在文化梳理的基础上，探索用AIGC工具生产傩戏短视频的具体方法：用文生图'
    '技术重构傩面具的视觉表达，让纹样"动起来、活起来"；用文生视频技术还原傩舞的程式动作，把"搬先锋"'
    '"开山猛将"等经典剧目搬进数字场景；用数字人技术生成傩戏解说与虚拟展示。重点解决"怎么把文化元素'
    '变成短视频"这一技术问题，形成一套可操作的生产流程和提示词模板。')
add_body('在创作过程中，坚持"AI生成+传承人把关"的双重机制：AI负责效率，传承人负责把关，确保每一个'
    '生成的傩面具、每一段还原的傩舞都不偏离文化本真。这一机制既是质量保障，也是项目区别于一般AI'
    '内容的特色所在。')

add_title('3. 傩戏短视频的多平台传播策略')
add_body('内容是"做出来"，传播是"送出去"。研究将把产出的傩戏短视频投放到抖音、B站、微信视频号等'
    '主流平台，观察不同平台的受众反馈，比较各平台的内容偏好和传播效果。重点研究三个问题：什么样的'
    '傩戏短视频更容易被年轻用户接受；怎样的标题、封面、节奏能提升完播率；如何通过系列化、IP化运营'
    '形成持续关注。用真实的传播数据来检验内容生产的有效性，而不是停留在"做完了"的层面。')

add_title('4. 傩戏数字资源库的沉淀与传承')
add_body('研究不是一次性的内容生产，而是要沉淀可复用的数字资产。项目将把过程中积累的傩戏素材、'
    'AIGC提示词模板、生成的角色与场景、传播数据等整理为结构化的数字资源库，形成"傩戏—AIGC"内容'
    '生产的标准操作流程。这套资源库既能让后续的内容生产"站在前人肩膀上"，也能作为池州傩戏数字化'
    '传承的基础设施，供地方文旅、教育机构持续使用。')

add_title('（二）项目拟解决的关键问题')
add_body('围绕上述研究内容，项目拟解决三个关键问题。')

add_title('1. 传统傩戏如何转化为适配短视频的内容形态')
add_body('傩戏的完整演出长达数小时，而短视频的黄金时长只有几十秒。这中间的转化不是简单"剪短"，'
    '而是要找到傩戏中最有辨识度、最适合碎片化传播的"记忆点"——一面狞厉的傩面具、一个刚劲的傩舞'
    '定格、一句有韵味的傩戏唱腔。研究要解决的第一个问题，就是把厚重的傩戏"翻译"成短视频语言，'
    '既不失其神，又能抓住眼球。')

add_title('2. AIGC生成如何守住傩戏文化本真')
add_body('这是技术伦理层面的关键问题。AIGC生成的傩面具可能"神似而形散"，傩舞动作可能"流畅而走样"。'
    '如果放任AI自由发挥，产出的内容就会变成"伪非遗"，反而伤害傩戏的文化严肃性。研究要通过"传承人'
    '把关、文化馆校准、纹样溯源"三重机制，为AI生成划定边界，回答"AI做非遗，哪些能变、哪些不能变"'
    '这一核心问题。')

add_title('3. 傩戏短视频如何触达年轻受众实现破圈')
add_body('傩戏年轻受众触达率低，是多年未解的老问题。已有的破圈尝试（如入驻《桃源深处有人家》手游、'
    '拍摄微短剧）提供了参考，但短视频这条路径还没被系统验证。研究要解决的第三个问题，就是通过真实的'
    '多平台投放和数据反馈，找到傩戏短视频"破圈"的有效打法，让古老的傩戏真正走进年轻人的手机屏幕。')

# ===== 三、项目研究与实施的基础条件 =====
add_title('三、项目研究与实施的基础条件', size=14)
add_body('本项目具备较好的实施基础。首先是地域优势：池州学院地处池州，傩戏的主要发源地——贵池区'
    '梅街镇刘街社区就在本地，项目团队可以方便地实地调研傩戏演出、走访傩文化展示馆、接触传承人和'
    '傩戏艺人，获取第一手文化素材，这是外地高校难以具备的条件。')
add_body('其次是前期积累。池州傩戏近年来已有数字化探索的初步实践：傩戏入驻腾讯手游《桃源深处有人家》'
    '，12尊傩面具角色完成数字化转化；首部傩文化微短剧《花裙彩面踏春芳》已拍摄完成；《解码非遗》系列'
    '短视频在线上发布；梅街镇12个村（社区）都建立了自己的宣传账号。这些实践为本项目提供了可借鉴的'
    '案例素材和内容基础。')
add_body('第三是技术条件。项目使用的AIGC工具（文生图、文生视频、数字人等）已较为成熟、使用门槛低，'
    '团队成员在课程学习和实践中已初步掌握AI绘图、智能剪辑等基本技能。学校的机房和网络条件能够支撑'
    'AI工具的稳定运行，指导教师可在技术方法和内容把关方面提供指导。')
add_body('需要说明的是，项目对傩戏文化内涵的理解可能还不够深入，AI生成内容的专业性也需要在实施中'
    '逐步提升。对此，项目将通过"传承人把关"机制来弥补，把文化准确性交给最懂傩戏的人，而不是单纯'
    '依赖技术。')

# ===== 四、项目实施方案 =====
add_title('四、项目实施方案', size=14)
add_title('第一阶段：研究深化与方案细化（2026年5月-2026年9月）')
add_body('系统调研池州傩戏文化资源，整理傩面具谱系与剧目资料，确定AIGC工具方案和内容生产规范。')
add_title('第二阶段：实地调研与数据整合（2026年10月-2027年4月）')
add_body('走访傩文化展示馆与傩戏发源地，采集傩面具、傩舞影像素材，建立傩戏文化素材库。')
add_title('第三阶段：模型构建与综合分析（2027年5月-2027年9月）')
add_body('完成傩戏AIGC短视频的内容生产与多平台投放，收集传播数据，优化内容与传播策略。')
add_title('第四阶段：成果凝练、转化与结题（2027年10月-2027年12月）')
add_body('沉淀傩戏数字资源库，撰写研究报告与论文，整理结题材料，推动成果向文旅应用转化。')

# ===== 五、需要学校提供的条件 =====
add_title('五、需要学校提供的条件', size=14)
add_sub('1.项目资金支持：', '学校提供资金支持，主要用于差旅（实地调研傩戏发源地）、图书资料购买、数据采集分析、成果出版打印等，以保障、推进项目的实施。')
add_sub('2.实验场地支持：', '池州学院配备机房以及工作室，可满足项目在AI内容生成、视频剪辑制作等方面的场地与设备需求。')

# ===== 六、预期成果 =====
add_title('六、预期成果', size=14)
add_sub('1. 学术论文：', '在省级及以上期刊发表AIGC与非遗数字化传播方向的论文1篇。')
add_sub('2. 短视频作品：', '产出池州傩戏AIGC主题短视频系列不少于10条，并在抖音、B站、微信视频号等平台发布。')
add_sub('3. 数字资源库：', '建成池州傩戏AIGC数字资源库1套，含傩面具素材、提示词模板与内容生产流程规范。')
add_body('项目后续预期成果：将项目形成的内容生产范式推广到青阳腔、东至花灯等池州其他非遗项目，为地方文旅宣传持续提供数字内容支撑。')

doc.save(OUTPUT)
print(f"完成: {OUTPUT}")

# 统计字数
doc2 = Document(OUTPUT)
c0 = doc2.tables[3].cell(0, 0)
total0 = sum(len(p.text) for p in c0.paragraphs)
c1 = doc2.tables[3].cell(1, 0)
total1 = sum(len(p.text) for p in c1.paragraphs)
print(f"表3行0(一~六章): ~{total0}字")
print(f"表3行1(经费): ~{total1}字")
print(f"正文总计: ~{total0+total1}字")
